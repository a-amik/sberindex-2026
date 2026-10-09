"""Таблицы Росстата.

* ИПЦ по субъектам к предыдущему месяцу: всего, продовольственные,
  непродовольственные товары, услуги; листы вида «03(2024)»;
* среднемесячная номинальная начисленная зарплата по субъектам;
* численность постоянного населения МО на 1 января (ТЕРСОН-МО);
* монопрофильные МО с категорией социально-экономического положения.

В имени файла стоит месяц публикации, поэтому ссылка ищется на странице
раздела, а не задаётся жёстко.
"""
from __future__ import annotations

import re
import urllib.parse

import openpyxl
import pandas as pd

from .. import net
from ..regions import key as region_key

MONTHS = ["январ", "феврал", "март", "апрел", "ма", "июн", "июл", "август", "сентябр", "октябр", "ноябр", "декабр"]


def download(cfg) -> None:
    raw = cfg.path("raw") / "rosstat"
    raw.mkdir(parents=True, exist_ok=True)
    base = cfg["sources"]["rosstat"]
    s = net.session()
    pages: dict[str, str] = {}
    for name, (page, pattern) in cfg["rosstat_files"].items():
        target = raw / f"{name}.xlsx"
        if target.exists():
            continue
        if page not in pages:
            pages[page] = net.fetch(s, base + page).text
        # имя файла бывает с кириллической буквой в начале: «Сhisl_MO…»
        hrefs = re.findall(r'href="(/storage/mediabank/[^"]+)"', pages[page])
        found = sorted({h for h in hrefs if re.search(pattern, h)})
        if not found:
            raise RuntimeError(f"Росстат: на {page} нет файла по шаблону {pattern}")
        url = base + urllib.parse.quote(found[-1])
        target.write_bytes(net.fetch(s, url).content)
        net.record(cfg.path("manifest"), target, url)
        print(f"  {name}: {found[-1].rsplit('/', 1)[-1]}")
    parse_all(cfg)


def _num(x):
    try:
        return float(str(x).replace(",", ".").replace("\xa0", "").strip())
    except ValueError:
        return float("nan")


def parse_cpi(path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = []
    for sh in wb.sheetnames:
        m = re.fullmatch(r"(\d{2})\((\d{4})\)", sh.strip())
        if not m:
            continue
        period = f"{m.group(2)}-{m.group(1)}"
        for r in wb[sh].iter_rows(values_only=True):
            code, name = r[0], r[1]
            if not isinstance(name, str) or not re.fullmatch(r"\d+", str(code or "").strip()):
                continue
            code = str(code).strip()
            if code == "643" or len(code) <= 2:
                kind = "rf" if code == "643" else "fd"
            else:
                kind = "region"
            v = [_num(x) for x in r[2:6]]
            rows.append({"period": period, "okato": code, "kind": kind, "region": name.strip(),
                         "cpi_total": v[0], "cpi_food": v[1], "cpi_nonfood": v[2], "cpi_services": v[3]})
    df = pd.DataFrame(rows)
    df["region_key"] = df.region.map(region_key)
    return df


def parse_wages(path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["с 2019"]
    rows = list(ws.iter_rows(values_only=True))
    year_row = next(i for i, r in enumerate(rows) if any(str(x).startswith("2019") for x in r if x))
    years, year = [], None
    for x in rows[year_row]:
        if x is not None:
            year = int(str(x)[:4])
        years.append(year)
    months = rows[year_row + 1]
    cols = []
    for j, mname in enumerate(months):
        if not mname:
            cols.append(None)
            continue
        idx = next((k for k, p in enumerate(MONTHS) if str(mname).strip().lower().startswith(p)), None)
        cols.append(f"{years[j]}-{idx + 1:02d}" if idx is not None and years[j] else None)
    out = []
    for r in rows[year_row + 2:]:
        name = r[0]
        if not isinstance(name, str) or not name.strip():
            continue
        for j, period in enumerate(cols):
            if period and j < len(r) and r[j] not in (None, "", "…", "-"):
                out.append({"period": period, "region": name.strip(), "wage": _num(r[j])})
    df = pd.DataFrame(out).dropna(subset=["wage"])
    df["region_key"] = df.region.map(region_key)
    return df


def parse_population(path) -> pd.DataFrame:
    """Численность на 1 января по МО: код ТЕРСОН-МО, всё, городское, сельское население."""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Численность_по_МО"]
    year = int(re.search(r"(\d{4})", path.name).group(1)) if re.search(r"\d{4}", path.name) else None
    out = []
    for r in ws.iter_rows(values_only=True):
        code = re.sub(r"\D", "", str(r[0] or ""))   # бывают пробелы внутри кода
        if re.fullmatch(r"\d{9}|\d{14}", code):   # Excel теряет ведущий ноль у кодов 01…
            code = "0" + code
        if not re.fullmatch(r"\d{10}", code):   # 10 знаков — МО; 15 — населённые пункты
            continue
        out.append({"terson": code, "oktmo8": code[:8], "name": str(r[1]).strip(),
                    "pop": _num(r[2]), "pop_urban": _num(r[3]), "pop_rural": _num(r[4])})
    return pd.DataFrame(out)


def parse_monotowns(path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True)
    out = []
    for sh, cat in (("Сложное", "сложное"), ("Риски", "риски"), ("Стабильное", "стабильное")):
        for r in wb[sh].iter_rows(values_only=True):
            if isinstance(r[0], (int, float)) or (isinstance(r[0], str) and r[0].strip().isdigit()):
                out.append({"region": str(r[1]).strip(), "mo_name": str(r[2]).strip(), "mono_status": cat})
    df = pd.DataFrame(out)
    df["region_key"] = df.region.map(region_key)
    return df


def parse_all(cfg) -> None:
    raw = cfg.path("raw") / "rosstat"
    out = cfg.path("external") / "rosstat"
    out.mkdir(parents=True, exist_ok=True)
    parse_cpi(raw / "cpi_regions.xlsx").to_parquet(out / "cpi_regions.parquet", index=False)
    parse_wages(raw / "wages_regions.xlsx").to_parquet(out / "wages_regions.parquet", index=False)
    pop = parse_population(raw / "population_mo.xlsx")
    pop["year"] = 2025
    pop.to_parquet(out / "population_mo.parquet", index=False)
    parse_monotowns(raw / "monotowns.xlsx").to_parquet(out / "monotowns.parquet", index=False)
